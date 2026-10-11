"""Presentations: AI writes the content, python-pptx lays it out in a calm style."""
from __future__ import annotations

import json
import re

from langchain_core.messages import HumanMessage, SystemMessage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from app.agent.llm import get_llm, has_ai
from app.agent.prompt_loader import load_prompt
from app.config import settings
from app.services import files

PINE = RGBColor(0x2E, 0x5E, 0x4E)
INK = RGBColor(0x1E, 0x2A, 0x23)
MUTED = RGBColor(0x5F, 0x6C, 0x65)
MIST = RGBColor(0xED, 0xF0, 0xEA)
POLLEN = RGBColor(0xF1, 0xE3, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def _fallback(topic: str, count: int) -> list[dict]:
    base = [
        ("Why this matters", [f"What {topic} is, in one sentence", "The problem it solves", "Who it's for"]),
        ("How it works", ["The main parts", "How they fit together", "A simple example"]),
        ("Where it's used", ["Everyday use cases", "Results people see", "Who's doing it well"]),
        ("What to watch out for", ["Common mistakes", "Costs and trade-offs", "How to reduce the risks"]),
        ("Getting started", ["A first small step", "Tools and resources", "How to measure progress"]),
        ("Key takeaways", [f"{topic} in a nutshell", "One thing to try this week", "Questions to explore next"]),
    ]
    return [{"title": t, "bullets": b} for t, b in base[: max(3, count)]]


def _parse(text: str) -> list[dict]:
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    start, end = text.find("["), text.rfind("]")
    data = json.loads(text[start : end + 1])
    slides = []
    for s in data:
        title = str(s.get("title", "")).strip()
        bullets = [str(b).strip() for b in s.get("bullets", []) if str(b).strip()][:6]
        if title and bullets:
            slides.append({"title": title[:90], "bullets": [b[:160] for b in bullets]})
    if len(slides) < 2:
        raise ValueError("too few slides")
    return slides


def write_slides(topic: str, count: int, notes: str = "") -> tuple[list[dict], bool]:
    """Returns (slides, written_by_ai)."""
    if has_ai():
        try:
            prompt = load_prompt("slides.txt").replace("{count}", str(count))
            ask = f"Topic: {topic}\nNumber of slides: {count}" + (f"\nExtra guidance from the user: {notes}" if notes else "")
            reply = get_llm(temperature=0.5).invoke([SystemMessage(content=prompt), HumanMessage(content=ask)])
            return _parse(str(reply.content)), True
        except Exception:  # noqa: BLE001 — fall back to a sensible outline
            pass
    return _fallback(topic, count), False


def _bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _text(slide, left, top, width, height, text, size, color, bold=False, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    return tf


def build(topic: str, slides: list[dict]) -> str:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    W, H = prs.slide_width, prs.slide_height
    blank = prs.slide_layouts[6]

    # Title slide
    s = prs.slides.add_slide(blank)
    _bg(s, PINE)
    _text(s, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(1.8), topic, 48, WHITE, bold=True)
    _text(s, Inches(0.9), Inches(4.3), W - Inches(1.8), Inches(0.6), f"Prepared with Ask {settings().user_name}", 18, POLLEN)
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, W - Inches(2.2), Inches(0.8), Inches(1.2), Inches(1.2))
    dot.fill.solid()
    dot.fill.fore_color.rgb = POLLEN
    dot.line.fill.background()

    for i, data in enumerate(slides, start=1):
        s = prs.slides.add_slide(blank)
        _bg(s, MIST)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(0.9), Inches(0.12), Inches(0.9))
        bar.fill.solid()
        bar.fill.fore_color.rgb = PINE
        bar.line.fill.background()
        _text(s, Inches(1.2), Inches(0.8), W - Inches(2.4), Inches(1.1), data["title"], 34, INK, bold=True)
        tf = _text(s, Inches(1.2), Inches(2.2), W - Inches(2.6), H - Inches(3.2), "", 22, INK)
        tf.paragraphs[0].text = ""
        for j, b in enumerate(data["bullets"]):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.size = Pt(22)
            p.font.color.rgb = INK
            p.space_after = Pt(14)
        _text(s, W - Inches(1.6), H - Inches(0.8), Inches(1), Inches(0.4), f"{i}", 12, MUTED)

    out = files.output_path(f"{files.safe_name(topic)[:60]} slides.pptx")
    prs.save(out)
    return out.name


def create(topic: str, count: int = 6, notes: str = "") -> dict:
    topic = " ".join((topic or "").split())[:100] or "My presentation"
    count = max(3, min(int(count or 6), 12))
    slides, by_ai = write_slides(topic, count, notes)
    name = build(topic, slides)
    return {"file": name, "slides": len(slides) + 1, "by_ai": by_ai, "titles": [s["title"] for s in slides]}
